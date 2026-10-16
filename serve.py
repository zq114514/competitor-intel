# -*- coding: utf-8 -*-
"""
竞品技术情报站 本地服务：静态站点 + /ai 大模型中转代理
用法：python serve.py   （替代 python -m http.server 8765）
- 静态文件服务 webapp/ 于 http://localhost:8765
- POST /ai  {agentId, content, sessionId} → 转发火山方舟 chat/completions（SSE流式）
  密钥只存于本机 ai_config.json，绝不进前端代码与git仓库
"""
import json, os, sys, re, glob, urllib.request, urllib.error
from http.server import ThreadingHTTPServer, SimpleHTTPRequestHandler
from urllib.parse import urlparse, parse_qs

BASE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.join(BASE, 'webapp')
CONF_PATH = os.path.join(BASE, 'ai_config.json')
ARK_URL = 'https://ark.cn-beijing.volces.com/api/v3/chat/completions'
PORT = 8765

# 图片上传：Content-Type → 扩展名
UPLOAD_EXT = {'image/jpeg': 'jpg', 'image/jpg': 'jpg', 'image/png': 'png',
              'image/webp': 'webp', 'image/gif': 'gif'}
UPLOAD_MAX = 20 * 1024 * 1024


def load_cfg():
    with open(CONF_PATH, encoding='utf-8') as f:
        return json.load(f)


class Server(ThreadingHTTPServer):
    # Windows默认SO_REUSEADDR允许多进程绑同一端口（请求被随机分发到僵尸实例），关闭之
    allow_reuse_address = False
    daemon_threads = True


class Handler(SimpleHTTPRequestHandler):
    def __init__(self, *a, **kw):
        super().__init__(*a, directory=ROOT, **kw)

    def end_headers(self):
        # 月份数据文件与图片经管理页上传后会变：禁用浏览器缓存，刷新即见最新
        p = self.path.split('?')[0]
        if p.startswith('/js/data') or p.startswith('/img/'):
            self.send_header('Cache-Control', 'no-cache, must-revalidate')
        super().end_headers()

    def log_message(self, fmt, *args):  # 静默静态资源日志
        if self.path == '/ai' or self.path.startswith('/ai?'):
            sys.stderr.write('[ai] %s\n' % (fmt % args))

    def do_POST(self):
        route = self.path.split('?')[0]
        if route == '/upload':
            self.handle_upload()
            return
        if route == '/build-shadow':
            self.handle_build_shadow()
            return
        if route != '/ai':
            self.send_error(404)
            return
        try:
            n = int(self.headers.get('Content-Length') or 0)
            body = json.loads(self.rfile.read(n) or b'{}')
        except Exception:
            self._json_err(400, '请求体不是合法JSON')
            return
        content = (body.get('content') or '').strip()
        if not content:
            self._json_err(400, 'content为空')
            return
        if len(content) > 30000:
            content = content[:30000]
        try:
            cfg = load_cfg()
        except Exception as e:
            self._json_err(500, 'ai_config.json读取失败: %s' % e)
            return
        model = body.get('agentId') or cfg.get('model')
        payload = json.dumps({
            'model': model,
            'messages': [
                {'role': 'system', 'content': cfg.get('system', '你是汽车整车性能集成专家，回答严谨、基于给定数据，不编造。')},
                {'role': 'user', 'content': content},
            ],
            'stream': True,
        }).encode('utf-8')
        req = urllib.request.Request(ARK_URL, data=payload, headers={
            'Content-Type': 'application/json',
            'Authorization': 'Bearer ' + cfg['api_key'],
        })
        try:
            up = urllib.request.urlopen(req, timeout=180)
        except urllib.error.HTTPError as e:
            detail = e.read()[:600].decode('utf-8', 'replace')
            self._json_err(e.code, '方舟API错误: %s' % detail)
            return
        except Exception as e:
            self._json_err(502, '方舟连接失败: %s' % e)
            return
        # 把方舟OpenAI格式SSE翻译成前端期望的 data:{"delta":"..."}
        self.send_response(200)
        self.send_header('Content-Type', 'text/event-stream; charset=utf-8')
        self.send_header('Cache-Control', 'no-cache')
        self.send_header('X-Accel-Buffering', 'no')
        self.end_headers()
        try:
            while True:
                line = up.readline()  # SSE按行推流，readline避免缓冲阻塞
                if not line:
                    break
                line = line.strip()
                if not line.startswith(b'data:'):
                    continue
                data = line[5:].strip()
                if data == b'[DONE]':
                    continue
                try:
                    j = json.loads(data)
                    delta = (j.get('choices') or [{}])[0].get('delta', {}).get('content')
                except Exception:
                    delta = None
                if delta:
                    out = 'data: %s\n\n' % json.dumps({'delta': delta}, ensure_ascii=False)
                    self.wfile.write(out.encode('utf-8'))
                    self.wfile.flush()
        except (BrokenPipeError, ConnectionResetError):
            pass
        finally:
            up.close()

    def handle_build_shadow(self):
        """POST /build-shadow  body=前端按当前口径导出的单车LACU结构（含各版型），
        用openpyxl生成 webapp/shadow/lacu/{month}_{id}.xlsx（WPS可直接打开）：
        说明页 + 每版型一页（L/A/C/U界限表 + 竞品明细表，本车金色行，来源可点）。"""
        try:
            n = int(self.headers.get('Content-Length') or 0)
            if n <= 0 or n > 20 * 1024 * 1024:
                self._json_err(400, '体积非法'); return
            payload = json.loads(self.rfile.read(n) or b'{}')
        except Exception:
            self._json_err(400, '请求体不是合法JSON'); return
        month = re.sub(r'[^0-9-]', '', str(payload.get('month') or ''))
        cid = re.sub(r'[^A-Za-z0-9_\-]', '', str(payload.get('id') or ''))
        if not re.match(r'^\d{4}-\d{2}$', month) or not cid:
            self._json_err(400, 'month/id非法'); return
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
            from openpyxl.utils import get_column_letter
        except Exception as e:
            self._json_err(500, 'openpyxl不可用: %s' % e); return

        NAVY = 'FF16233F'; GOLD = 'FFE8B95C'; GREY = 'FFE9EDF5'; RED = 'FFF04B4C'
        thin = Side(style='thin', color='FFB8C2D6')
        border = Border(left=thin, right=thin, top=thin, bottom=thin)
        hdr_font = Font(color='FFFFFFFF', bold=True, size=10)
        hdr_fill = PatternFill('solid', fgColor=NAVY)
        gold_fill = PatternFill('solid', fgColor='33E8B95C')
        self_fill = PatternFill('solid', fgColor=GOLD)
        center = Alignment(horizontal='center', vertical='center', wrap_text=True)
        left = Alignment(horizontal='left', vertical='center', wrap_text=True)

        wb = Workbook()
        ws0 = wb.active
        ws0.title = '说明'
        info = [
            ['LACU 竞品数据库 · 单车竞品工作簿'],
            ['车型', payload.get('name', cid)],
            ['数据基准月', payload.get('monthLabel', month)],
            ['生成时间', payload.get('builtAt', '')],
            ['圈选口径', '同级×同车身×指导价带（±max(5万,本车价15%)），竞品指导价区间与价格带相交即入圈；同级有价竞品不足4款时并入相邻级别车型（明细以"邻"标注，档级仅供参考）；续航/补能按同能源组排名'],
            ['档位定义', 'L前10%领先/A前30%/C中位/U后30%；n<4不出档，n=4-6仅L/C/U'],
            ['操稳稳定口径', '四因子加权：最小转弯半径50%（组内百分位，越小越优）+前悬结构档20%+后悬结构档20%+可调悬架折分10%（无0/软硬1/软硬+高低2），缺项按有值项权重重归一；结构档：前悬双叉臂/双横臂/多连杆=3、麦弗逊/连杆支柱=2，后悬多连杆/双叉臂=3、双横臂/连杆支柱/拖臂=2、纵臂扭转梁=1，整体桥硬派越野（带低速四驱）记3、城市车记1；结构分级不含调校水平；空悬/魔毯/电磁/后轮转向等配置平铺不计分'],
            ['数据来源', '懂车帝/汽车之家配置表、车企官网及上市通稿；明细含逐车链接可核对，无官方数据不推算'],
        ]
        for r in info:
            ws0.append(r)
        ws0['A1'].font = Font(bold=True, size=14)
        ws0.merge_cells('A1:B1')
        for row in ws0.iter_rows(min_row=2, max_row=ws0.max_row, max_col=2):
            row[0].font = Font(bold=True)
            row[0].fill = PatternFill('solid', fgColor=GREY)
            for c in row:
                c.alignment = left; c.border = border
        ws0.column_dimensions['A'].width = 14
        ws0.column_dimensions['B'].width = 110
        for i in range(5, 10):
            ws0.row_dimensions[i].height = 42
        ws0.row_dimensions[7].height = 120

        AX_HDR = ['维度', '本车值', '名次', 'n', '档位', 'L界限(前10%)', 'A界限(前30%)', 'C中位', 'U界限(后30%)']
        MEM_HDR = ['名次', '车型', '指导价(万)', '零百(s)', '纯电续航(km)', '综合续航(km)', '等效快充20-80%(min)', '轴距(mm)',
                   '转弯半径(m)', '前/后悬(结构档)', '可调悬架', '驱动方式', '参评维度', '来源']

        used_titles = set()
        def safe_title(s, idx):
            t = re.sub(r'[\\/?*\[\]:]', '', str(s))[:28] or '口径'
            if t in used_titles:
                suf = '#' + str(idx)
                t = t[:28 - len(suf)] + suf
            used_titles.add(t)
            return t

        for ti, tab in enumerate(payload.get('tabs', [])):
            ws = wb.create_sheet(safe_title(tab.get('label'), ti))
            ws.append([payload.get('name', cid) + '  ·  ' + str(tab.get('label')) + '  ·  ' + re.sub(r'<[^>]+>', '', str(tab.get('band', '')))])
            ws.merge_cells('A1:N1')
            ws['A1'].font = Font(bold=True, size=11, color='FFB8860B')
            ws['A1'].alignment = left
            ws.append(AX_HDR)
            for c in ws[2]:
                c.font = hdr_font; c.fill = hdr_fill; c.alignment = center; c.border = border
            for ax in tab.get('axes', []):
                ws.append([ax.get('name'), ax.get('selfV'), ax.get('selfRank'), ax.get('n'),
                           ax.get('selfBand'), ax.get('L'), ax.get('A'), ax.get('C'), ax.get('U')])
                rr = ws.max_row
                for ci, c in enumerate(ws[rr], start=1):
                    c.alignment = center; c.border = border
                    if ci == 5 and ax.get('selfBand'):
                        v = ax.get('selfBand')
                        c.fill = PatternFill('solid', fgColor={'L': 'FF2E8B57', 'A': 'FF4A90D9', 'C': 'FF8A94A6', 'U': 'FFE05B52'}.get(v, GREY))
                        c.font = Font(color='FFFFFFFF', bold=True)
            end_ax = ws.max_row
            ws.append([])
            ws.append(MEM_HDR)
            hr = ws.max_row
            for c in ws[hr]:
                c.font = hdr_font; c.fill = hdr_fill; c.alignment = center; c.border = border
            for m in tab.get('members', []):
                ws.append([m.get('rank'), m.get('name'), m.get('price'), m.get('accel'),
                           m.get('rangeKm'), m.get('totalRangeKm'), m.get('eqMin'), m.get('wheelbase'),
                           m.get('turn'), m.get('susp'), m.get('suspAdj'), m.get('drive'),
                           m.get('axes'), m.get('src')])
                rr = ws.max_row
                is_self = bool(m.get('self'))
                for ci, c in enumerate(ws[rr], start=1):
                    c.alignment = center if ci not in (2, 10, 11, 12, 13, 14) else left
                    c.border = border
                    if is_self:
                        c.fill = self_fill; c.font = Font(bold=True)
                if m.get('url'):
                    cell = ws.cell(row=rr, column=14)
                    cell.hyperlink = m['url']
                    cell.font = Font(color='FF1A5FB4', underline='single', bold=is_self)
            widths = [6, 30, 10, 8, 12, 13, 16, 10, 9, 16, 11, 15, 22, 20]
            for i, w in enumerate(widths, start=1):
                ws.column_dimensions[get_column_letter(i)].width = w
            ws.freeze_panes = 'A3'

        outdir = os.path.join(ROOT, 'shadow', 'lacu')
        os.makedirs(outdir, exist_ok=True)
        fpath = os.path.join(outdir, '%s_%s.xlsx' % (month, cid))
        wb.save(fpath)
        b = json.dumps({'ok': True, 'path': 'shadow/lacu/%s_%s.xlsx' % (month, cid)},
                       ensure_ascii=False).encode('utf-8')
        self.send_response(200)
        self.send_header('Content-Type', 'application/json; charset=utf-8')
        self.send_header('Cache-Control', 'no-store')
        self.send_header('Content-Length', str(len(b)))
        self.end_headers()
        self.wfile.write(b)

    def handle_upload(self):
        """POST /upload?cid=车型ID&file=js/data.js ；body=图片原始二进制。
        自动按 {cid}.{ext} 存入 webapp/img/，删除同cid旧扩展名文件，并改写数据文件img字段。"""
        q = parse_qs(urlparse(self.path).query)
        cid = (q.get('cid') or [''])[0].strip()
        rel = (q.get('file') or ['js/data.js'])[0].strip()
        if not re.fullmatch(r'[a-zA-Z0-9_\-]+', cid):
            self._json_err(400, 'cid非法'); return
        base = os.path.basename(rel.replace('\\', '/'))  # 防目录穿越，只取文件名
        if not re.fullmatch(r'data[\w\-]*\.js', base):
            self._json_err(400, 'file参数非法（仅允许js/data*.js）'); return
        data_path = os.path.join(ROOT, 'js', base)
        if not os.path.exists(data_path):
            self._json_err(404, '数据文件不存在: %s' % base); return
        ctype = (self.headers.get('Content-Type') or '').split(';')[0].strip().lower()
        ext = UPLOAD_EXT.get(ctype)
        if not ext:
            self._json_err(400, '仅支持 JPG/PNG/WebP/GIF（当前Content-Type: %s）' % ctype); return
        n = int(self.headers.get('Content-Length') or 0)
        if n <= 0 or n > UPLOAD_MAX:
            self._json_err(400, '文件为空或超过20MB'); return
        blob = self.rfile.read(n)
        if len(blob) != n:
            self._json_err(400, '上传数据不完整'); return
        # 存图（先写临时名再落位，避免半文件）
        img_dir = os.path.join(ROOT, 'img')
        os.makedirs(img_dir, exist_ok=True)
        fname = '%s.%s' % (cid, ext)
        tmp_path = os.path.join(img_dir, fname + '.part')
        with open(tmp_path, 'wb') as f:
            f.write(blob)
        # 清理同cid其他扩展名旧文件（中文名旧文件不在此规则内，保留）
        for old in glob.glob(os.path.join(img_dir, cid + '.*')):
            ob = os.path.basename(old)
            if ob != fname and not ob.endswith('.part'):
                try: os.remove(old)
                except OSError: pass
        os.replace(tmp_path, os.path.join(img_dir, fname))
        # 改写数据文件中该车的img字段（从 id:'cid' 锚点到其块内第一个 img: '...'，lambda替换防转义坑）
        with open(data_path, encoding='utf-8') as f:
            src = f.read()
        pat = re.compile(r"(id:\s*'" + re.escape(cid) + r"'[\s\S]*?img:\s*')[^']*(')")
        new_field = 'img/' + fname
        new_src, cnt = pat.subn(lambda m: m.group(1) + new_field + m.group(2), src, count=1)
        if cnt == 0:
            self._json_err(500, "数据文件 %s 中未找到车型 %s（图片已保存，img字段需手工补）" % (base, cid)); return
        with open(data_path, 'w', encoding='utf-8', newline='') as f:
            f.write(new_src)
        b = json.dumps({'ok': True, 'path': new_field, 'file': 'js/' + base}, ensure_ascii=False).encode('utf-8')
        self.send_response(200)
        self.send_header('Content-Type', 'application/json; charset=utf-8')
        self.send_header('Cache-Control', 'no-store')
        self.send_header('Content-Length', str(len(b)))
        self.end_headers()
        self.wfile.write(b)

    def _json_err(self, code, msg):
        b = json.dumps({'error': msg}, ensure_ascii=False).encode('utf-8')
        self.send_response(code)
        self.send_header('Content-Type', 'application/json; charset=utf-8')
        self.send_header('Content-Length', str(len(b)))
        self.end_headers()
        self.wfile.write(b)


if __name__ == '__main__':
    if not os.path.exists(CONF_PATH):
        print('缺少 ai_config.json（含api_key/model），请先创建')
        sys.exit(1)
    cfg = load_cfg()
    try:
        srv = Server(('127.0.0.1', PORT), Handler)
    except OSError:
        print('端口 %d 已被占用——服务可能已在运行，请勿重复启动。' % PORT)
        sys.exit(0)
    print('竞品情报站 + AI代理 已启动: http://localhost:%d  (模型: %s)' % (PORT, cfg.get('model')))
    srv.serve_forever()
